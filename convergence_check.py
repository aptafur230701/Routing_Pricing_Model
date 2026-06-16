"""
convergence_check.py
=====================
Chequeo automático de convergencia de entrenamiento PPO sobre training_log.

Evalúa 4 condiciones sobre una ventana móvil de las últimas W updates:
  1. Reward de entrenamiento aplanado (pendiente normalizada ≈ 0)
  2. Explained variance del crítico estable y >= umbral (no solo alta, sino sin tendencia)
  3. KL divergence consistentemente bajo el umbral de confianza
  4. Entropía estabilizada (pendiente ≈ 0, sin importar el nivel absoluto)

El modelo se considera "convergido" solo si las 4 condiciones se cumplen
simultáneamente. No usar entropía baja como proxy de convergencia: una
pendiente de entropía todavía negativa indica que la política sigue
reduciendo exploración activamente, sea o no su valor absoluto bajo.
"""

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


DEFAULT_THRESHOLDS = {
    'window':                       50,      # W: tamaño de la ventana móvil (últimas W updates)
    'reward_slope_rel_max':         0.001,   # |pendiente reward| / media reward en ventana
    'explained_var_min':            0.5,     # mismo umbral ya usado en plot_ppo_diagnostics
    'explained_var_slope_rel_max':  0.001,
    'kl_threshold':                 0.02,    # mismo umbral ya usado en plot_ppo_diagnostics
    'kl_violation_frac_max':        0.05,    # máx. fracción de updates en ventana que exceden kl_threshold
    'entropy_slope_abs_max':        0.0005,  # pendiente absoluta de entropía (no relativa, puede pasar por 0)
}


def _windowed_slope(values: np.ndarray) -> float:
    """Pendiente de regresión lineal simple (y = a*x + b) sobre el índice de la ventana."""
    n = len(values)
    if n < 2:
        return np.nan
    x = np.arange(n)
    slope, _ = np.polyfit(x, values, 1)
    return float(slope)


def check_convergence(training_log: list,
                      thresholds: dict = DEFAULT_THRESHOLDS) -> dict:
    """Evalúa las 4 condiciones de convergencia sobre la ventana final de training_log.

    Returns
    -------
    dict con:
      - 'converged': bool global
      - 'window_size_used': int (puede ser < thresholds['window'] si training_log es corto)
      - 'total_updates': int
      - 'conditions': dict con el detalle de cada condición (valor medido, umbral, pass/fail)
    """
    df = pd.DataFrame(training_log)
    required_cols = {'reward', 'entropy', 'kl_divergence', 'explained_var'}
    missing = required_cols - set(df.columns)
    if missing:
        raise KeyError(f"training_log no tiene las columnas requeridas: {missing}")

    W = min(thresholds['window'], len(df))
    if W < 2:
        raise ValueError(
            f"training_log tiene solo {len(df)} updates; "
            "se necesitan al menos 2 para evaluar pendiente."
        )

    window_df = df.tail(W)

    # Condición 1: reward aplanado
    reward_slope = _windowed_slope(window_df['reward'].to_numpy())
    reward_mean  = window_df['reward'].mean()
    reward_slope_rel = abs(reward_slope) / abs(reward_mean) if reward_mean != 0 else np.inf
    cond1_pass = reward_slope_rel < thresholds['reward_slope_rel_max']

    # Condición 2: explained variance estable y >= umbral
    ev_slope = _windowed_slope(window_df['explained_var'].to_numpy())
    ev_mean  = window_df['explained_var'].mean()
    ev_slope_rel = abs(ev_slope) / abs(ev_mean) if ev_mean != 0 else np.inf
    cond2_pass = (ev_mean >= thresholds['explained_var_min']) and \
                 (ev_slope_rel < thresholds['explained_var_slope_rel_max'])

    # Condición 3: KL divergence bajo control
    kl_violation_frac = (window_df['kl_divergence'] > thresholds['kl_threshold']).mean()
    cond3_pass = kl_violation_frac <= thresholds['kl_violation_frac_max']

    # Condición 4: entropía estabilizada (pendiente absoluta, no relativa)
    entropy_slope = _windowed_slope(window_df['entropy'].to_numpy())
    cond4_pass = abs(entropy_slope) < thresholds['entropy_slope_abs_max']

    conditions = {
        '1_reward_flattened': {
            'measured_slope':     reward_slope,
            'measured_slope_rel': reward_slope_rel,
            'threshold_rel':      thresholds['reward_slope_rel_max'],
            'window_mean':        reward_mean,
            'pass':               bool(cond1_pass),
        },
        '2_explained_var_stable': {
            'measured_mean':       ev_mean,
            'measured_slope_rel':  ev_slope_rel,
            'threshold_min':       thresholds['explained_var_min'],
            'threshold_slope_rel': thresholds['explained_var_slope_rel_max'],
            'pass':                bool(cond2_pass),
        },
        '3_kl_under_control': {
            'violation_fraction': float(kl_violation_frac),
            'threshold_fraction': thresholds['kl_violation_frac_max'],
            'kl_threshold':       thresholds['kl_threshold'],
            'pass':               bool(cond3_pass),
        },
        '4_entropy_stabilized': {
            'measured_slope': entropy_slope,
            'threshold_abs':  thresholds['entropy_slope_abs_max'],
            'window_mean':    float(window_df['entropy'].mean()),
            'pass':           bool(cond4_pass),
        },
    }

    converged = all(c['pass'] for c in conditions.values())

    return {
        'converged':        bool(converged),
        'window_size_used': W,
        'total_updates':    len(df),
        'conditions':       conditions,
    }


def print_convergence_report(result: dict) -> None:
    """Imprime un reporte legible del resultado de check_convergence."""
    print("=" * 65)
    print(f"  CONVERGENCE CHECK — window: last {result['window_size_used']} "
          f"of {result['total_updates']} updates")
    print("=" * 65)
    for name, c in result['conditions'].items():
        status = "PASS" if c['pass'] else "FAIL"
        print(f"\n  [{status}] {name}")
        for k, v in c.items():
            if k == 'pass':
                continue
            print(f"      {k}: {v:.6f}" if isinstance(v, float) else f"      {k}: {v}")
    print("\n" + "-" * 65)
    verdict = "CONVERGED" if result['converged'] else "NOT CONVERGED — extend training"
    print(f"  VEREDICTO: {verdict}")
    print("=" * 65)


# ──────────────────────────────────────────────────────────────────────────────
#  Excel export
# ──────────────────────────────────────────────────────────────────────────────

_ORANGE_DARK   = "C55A11"   # section headers
_ORANGE_LIGHT  = "FCE4D6"   # criteria table fill
_GRAY_HEADER   = "404040"   # data table column headers
_GRAY_LIGHT    = "D9D9D9"   # fixed columns
_GREEN_PASS    = "E2EFDA"   # PASS cells
_GREEN_DARK    = "548235"
_RED_FAIL      = "FFE0E0"   # FAIL cells
_RED_DARK      = "C00000"
_WHITE         = "FFFFFF"
_YELLOW        = "FFF2CC"   # measured values


def _xfill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _xfont(bold=False, color="000000", size=11):
    return Font(bold=bold, color=color, size=size)

def _xalign(h="center", wrap=True):
    return Alignment(horizontal=h, vertical="center", wrap_text=wrap)

def _xborder():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


def _sec_header(ws, row, col_start, col_end, text, fill_hex):
    ws.merge_cells(start_row=row, start_column=col_start,
                   end_row=row, end_column=col_end)
    c = ws.cell(row, col_start)
    c.value     = text
    c.fill      = _xfill(fill_hex)
    c.font      = _xfont(bold=True, color="FFFFFF", size=12)
    c.alignment = _xalign()
    c.border    = _xborder()


def _col_header(ws, row, col, text, fill_hex=_GRAY_HEADER, fcolor="FFFFFF"):
    c = ws.cell(row, col)
    c.value     = text
    c.fill      = _xfill(fill_hex)
    c.font      = _xfont(bold=True, color=fcolor)
    c.alignment = _xalign(wrap=True)
    c.border    = _xborder()


def _val_cell(ws, row, col, value, fill_hex=_WHITE, num_format=None, bold=False):
    c = ws.cell(row, col)
    c.value     = value
    c.fill      = _xfill(fill_hex)
    c.font      = _xfont(bold=bold)
    c.alignment = _xalign(h="center", wrap=False)
    c.border    = _xborder()
    if num_format and value is not None:
        c.number_format = num_format


def _result_cell(ws, row, col, passed: bool):
    text = "PASS" if passed else "FAIL"
    fill = _GREEN_PASS if passed else _RED_FAIL
    fc   = _GREEN_DARK if passed else _RED_DARK
    c = ws.cell(row, col)
    c.value     = text
    c.fill      = _xfill(fill)
    c.font      = _xfont(bold=True, color=fc)
    c.alignment = _xalign()
    c.border    = _xborder()


def save_convergence_to_excel(
    conv_result: dict,
    num_nodes: int,
    excel_path: str,
    thresholds: dict = DEFAULT_THRESHOLDS,
) -> None:
    """Añade (o reemplaza) la hoja 'PPO Convergence' en excel_path.

    El archivo debe existir ya (se llama después de build_formatted_excel).
    La hoja tiene dos secciones:
      · Criterios de convergencia — descripción explícita de cada condición y umbral
      · Resultados del chequeo   — tabla con valores medidos y PASS/FAIL por condición
    """
    wb = load_workbook(excel_path)
    if "PPO Convergence" in wb.sheetnames:
        del wb["PPO Convergence"]
    ws = wb.create_sheet("PPO Convergence")

    W   = thresholds['window']
    c   = conv_result['conditions']
    row = 1

    # ── SECCIÓN 1: Título ─────────────────────────────────────────────────────
    _sec_header(ws, row, 1, 9,
                f"CHEQUEO DE CONVERGENCIA PPO  ·  {num_nodes} nodos", _ORANGE_DARK)
    row += 2

    # ── SECCIÓN 2: Criterios ──────────────────────────────────────────────────
    _sec_header(ws, row, 1, 9,
                f"CRITERIOS DE CONVERGENCIA  (evaluados sobre las últimas W = {W} PPO updates)",
                _ORANGE_DARK)
    row += 1

    for col_idx, hdr in enumerate(
        ["Condición", "Métrica observada", "Cómo se mide", "Umbral de aprobación",
         "Nota de diseño"],
        1,
    ):
        _col_header(ws, row, col_idx, hdr, fill_hex="595959", fcolor="FFFFFF")
    row += 1

    criteria_table = [
        (
            "C1 — Reward aplanado",
            "Pendiente relativa del reward de entrenamiento",
            "Regresión lineal sobre la ventana → slope / |media reward|",
            f"< {thresholds['reward_slope_rel_max']}  (adimensional)",
            "Valor relativo para que sea comparable entre escalas de reward distintas.",
        ),
        (
            "C2 — Explained Variance estable",
            "Media y pendiente relativa de explained_var del crítico",
            "Media aritmética + regresión lineal sobre la ventana",
            f"Media ≥ {thresholds['explained_var_min']}  Y  |pendiente rel.| < {thresholds['explained_var_slope_rel_max']}",
            "Umbral 0.5 coincide con la línea de referencia en PPO_Diagnostics_AM.png.",
        ),
        (
            "C3 — KL divergence controlada",
            "Fracción de updates en ventana con KL > umbral KL",
            "Proporción de filas donde kl_divergence > kl_threshold",
            f"Fracción de violaciones ≤ {thresholds['kl_violation_frac_max'] * 100:.0f}%  "
            f"(umbral KL = {thresholds['kl_threshold']})",
            "Umbral KL 0.02 coincide con la línea de referencia en PPO_Diagnostics_AM.png.",
        ),
        (
            "C4 — Entropía estabilizada",
            "Pendiente ABSOLUTA de la entropía de la política",
            "Regresión lineal sobre la ventana → |slope| en nats/update",
            f"< {thresholds['entropy_slope_abs_max']}  nats/update",
            "Se usa pendiente absoluta (no relativa) porque la entropía puede pasar por 0.",
        ),
    ]

    for cond, metrica, como, umbral, nota in criteria_table:
        for col_idx, val in enumerate([cond, metrica, como, umbral, nota], 1):
            c_cell = ws.cell(row, col_idx)
            c_cell.value     = val
            c_cell.fill      = _xfill(_ORANGE_LIGHT)
            c_cell.font      = _xfont(bold=(col_idx == 1))
            c_cell.alignment = _xalign(h="left", wrap=True)
            c_cell.border    = _xborder()
        row += 1

    row += 2  # blank rows before results section

    # ── SECCIÓN 3: Resultados ─────────────────────────────────────────────────
    _sec_header(ws, row, 1, 17, "RESULTADOS DEL CHEQUEO", _ORANGE_DARK)
    row += 1

    # Sub-headers (merged groups)
    groups = [
        (1,  3,  "Identificación",          _GRAY_HEADER),
        (4,  6,  "C1 — Reward Aplanado",    "2E5395"),
        (7,  10, "C2 — Explained Variance", "7030A0"),
        (11, 13, "C3 — KL Divergence",      "548235"),
        (14, 16, "C4 — Entropía",           "833C00"),
        (17, 17, "Veredicto Final",         _ORANGE_DARK),
    ]
    for cs, ce, label, fill in groups:
        if cs == ce:
            ws.cell(row, cs).value = label
        else:
            ws.merge_cells(start_row=row, start_column=cs,
                           end_row=row, end_column=ce)
        c_cell = ws.cell(row, cs)
        c_cell.value     = label
        c_cell.fill      = _xfill(fill)
        c_cell.font      = _xfont(bold=True, color="FFFFFF")
        c_cell.alignment = _xalign()
        c_cell.border    = _xborder()
        for extra_col in range(cs + 1, ce + 1):
            ec = ws.cell(row, extra_col)
            ec.fill   = _xfill(fill)
            ec.border = _xborder()
    row += 1

    # Column headers
    col_headers = [
        # Identificación
        "Nodos",
        "Total\nUpdates",
        "Ventana\n(W usado)",
        # C1
        "Pendiente\nReward Rel.\n(medido)",
        f"Umbral\n(< {thresholds['reward_slope_rel_max']})",
        "C1\nResultado",
        # C2
        "EV Media\n(medido)",
        "EV Pendiente\nRel. (medido)",
        f"Umbral\nMedia ≥ {thresholds['explained_var_min']}",
        "C2\nResultado",
        # C3
        "Fracc.\nViolaciones\nKL (medido)",
        f"Umbral\n(≤ {thresholds['kl_violation_frac_max'] * 100:.0f}%)",
        "C3\nResultado",
        # C4
        "Pendiente\nEntropía Abs.\n(medido)",
        f"Umbral\n(< {thresholds['entropy_slope_abs_max']})",
        "C4\nResultado",
        # Veredicto
        "CONVERGIDO",
    ]
    group_fills = (
        [_GRAY_HEADER] * 3 +
        ["2E5395"] * 3 +
        ["7030A0"] * 4 +
        ["548235"] * 3 +
        ["833C00"] * 3 +
        [_ORANGE_DARK]
    )
    for col_idx, (hdr, fill) in enumerate(zip(col_headers, group_fills), 1):
        _col_header(ws, row, col_idx, hdr, fill_hex=fill)
    row += 1

    # ── Data row ──────────────────────────────────────────────────────────────
    cond   = conv_result['conditions']
    c1     = cond['1_reward_flattened']
    c2     = cond['2_explained_var_stable']
    c3     = cond['3_kl_under_control']
    c4     = cond['4_entropy_stabilized']

    _val_cell(ws, row, 1,  num_nodes,                              fill_hex=_GRAY_LIGHT, bold=True)
    _val_cell(ws, row, 2,  conv_result['total_updates'],           fill_hex=_GRAY_LIGHT)
    _val_cell(ws, row, 3,  conv_result['window_size_used'],        fill_hex=_GRAY_LIGHT)

    # C1
    _val_cell(ws, row, 4,  round(c1['measured_slope_rel'], 6),     fill_hex=_YELLOW, num_format="0.000000")
    _val_cell(ws, row, 5,  thresholds['reward_slope_rel_max'],     num_format="0.000000")
    _result_cell(ws, row, 6, c1['pass'])

    # C2
    _val_cell(ws, row, 7,  round(c2['measured_mean'], 4),          fill_hex=_YELLOW, num_format="0.0000")
    _val_cell(ws, row, 8,  round(c2['measured_slope_rel'], 6),     fill_hex=_YELLOW, num_format="0.000000")
    _val_cell(ws, row, 9,  thresholds['explained_var_min'],        num_format="0.00")
    _result_cell(ws, row, 10, c2['pass'])

    # C3
    _val_cell(ws, row, 11, round(c3['violation_fraction'], 4),     fill_hex=_YELLOW, num_format="0.0000")
    _val_cell(ws, row, 12, thresholds['kl_violation_frac_max'],    num_format="0.00")
    _result_cell(ws, row, 13, c3['pass'])

    # C4
    _val_cell(ws, row, 14, round(abs(c4['measured_slope']), 6),    fill_hex=_YELLOW, num_format="0.000000")
    _val_cell(ws, row, 15, thresholds['entropy_slope_abs_max'],    num_format="0.000000")
    _result_cell(ws, row, 16, c4['pass'])

    # Veredicto
    verdict_text = "SÍ" if conv_result['converged'] else "NO — extender entrenamiento"
    verdict_fill = _GREEN_PASS if conv_result['converged'] else _RED_FAIL
    verdict_fc   = _GREEN_DARK if conv_result['converged'] else _RED_DARK
    vc = ws.cell(row, 17)
    vc.value     = verdict_text
    vc.fill      = _xfill(verdict_fill)
    vc.font      = _xfont(bold=True, color=verdict_fc, size=12)
    vc.alignment = _xalign()
    vc.border    = _xborder()

    # ── Column widths ─────────────────────────────────────────────────────────
    widths = {
        1: 8,   2: 10,  3: 10,               # Identificación
        4: 16,  5: 12,  6: 10,               # C1
        7: 12,  8: 16,  9: 14,  10: 10,      # C2
        11: 14, 12: 12, 13: 10,              # C3
        14: 16, 15: 12, 16: 10,             # C4
        17: 24,                              # Veredicto
    }
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # Row heights
    ws.row_dimensions[row - 1].height = 48   # column headers (multi-line)
    for r in range(row - len(criteria_table) - 3, row - 3):
        ws.row_dimensions[r].height = 40

    ws.freeze_panes = "D6"

    wb.save(excel_path)
    print(f"Hoja 'PPO Convergence' guardada en: {excel_path}")


if __name__ == "__main__":
    import sys
    import pickle

    if len(sys.argv) < 2:
        print("Uso: python convergence_check.py <ruta_a_training_log.pkl>")
        sys.exit(1)

    with open(sys.argv[1], "rb") as f:
        training_log = pickle.load(f)

    result = check_convergence(training_log)
    print_convergence_report(result)
