"""
excel_report.py
===============
Builds a formatted Excel workbook with two sheets:
  · Summary         — 3 stacked blocks (det / stoch / timing)
  · Per Node Results — per-row solver data with per-method gap columns
"""

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Palette ───────────────────────────────────────────────────
_BLUE_DARK    = "2E5395"
_BLUE_LIGHT   = "DCE6F1"
_PURPLE_DARK  = "7030A0"
_PURPLE_LIGHT = "EBDDF0"
_GREEN_DARK   = "548235"
_GREEN_LIGHT  = "E2EFDA"
_YELLOW       = "FFF2CC"
_WHITE        = "FFFFFF"


# ── Style helpers ─────────────────────────────────────────────

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, color="000000", size=11):
    return Font(bold=bold, color=color, size=size)

def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def _border():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)

def _clean(v):
    """Convert nan/inf to None for Excel cells."""
    if v is None:
        return None
    try:
        f = float(v)
        if np.isnan(f) or np.isinf(f):
            return None
        return f
    except (TypeError, ValueError):
        return v

def _header_cell(cell, text, fill_hex, font_color="FFFFFF", bold=True):
    cell.value = text
    cell.fill = _fill(fill_hex)
    cell.font = _font(bold=bold, color=font_color)
    cell.alignment = _center()
    cell.border = _border()

def _data_cell(cell, value, num_format=None, fill_hex=None):
    cell.value = value
    cell.alignment = _left()
    cell.border = _border()
    if fill_hex:
        cell.fill = _fill(fill_hex)
    if value is not None and num_format:
        cell.number_format = num_format


# ── Gap calculation ───────────────────────────────────────────

def _compute_gap(benchmark, value):
    """Returns (benchmark - value) / abs(benchmark) * 100, or nan."""
    if benchmark is None or value is None:
        return float("nan")
    b, v = float(benchmark), float(value)
    if any(np.isnan(x) or np.isinf(x) for x in (b, v)) or b == 0:
        return float("nan")
    return (b - v) / abs(b) * 100

def _avg_gap(df, bench_col, bench_valid, method_col, method_valid):
    mask = df[bench_valid] & df[method_valid]
    if not mask.any():
        return float("nan")
    gaps = df.loc[mask].apply(
        lambda r: _compute_gap(r[bench_col], r[method_col]), axis=1
    ).dropna()
    return float(gaps.mean()) if len(gaps) > 0 else float("nan")


# ── Public entry point ────────────────────────────────────────

def build_formatted_excel(
    results_df: pd.DataFrame,
    summary_rows: list,
    output_path: str,
    node_size: int,
) -> None:
    wb = Workbook()
    _build_summary(wb, results_df, summary_rows)
    _build_per_node(wb, results_df)
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    wb.save(output_path)


# ═════════════════════════════════════════════════════════════
#  SUMMARY SHEET
# ═════════════════════════════════════════════════════════════

def _build_summary(wb, results_df, summary_rows):
    ws = wb.create_sheet("Summary")
    sr = summary_rows[0] if summary_rows else {}

    def g(key):
        return sr.get(key, float("nan"))

    # ── Precompute gaps ───────────────────────────────────────
    # Block 1 gaps vs LS-Exact
    gap_ls_exact     = 0.0
    gap_drl_det      = _clean(g("DRL Det Gap vs LS-Exact (%)"))
    gap_hga_lns      = _clean(_avg_gap(results_df, "LS-Exact Reward", "LS-Exact Valid",
                                        "HGA-LNS Reward",      "HGA-LNS Valid"))
    gap_rh_lookahead = _clean(_avg_gap(results_df, "LS-Exact Reward", "LS-Exact Valid",
                                        "RH-Lookahead Reward", "RH-Lookahead Valid"))
    gap_rh_greedy    = _clean(_avg_gap(results_df, "LS-Exact Reward", "LS-Exact Valid",
                                        "RH-Greedy Reward",    "RH-Greedy Valid"))

    # Block 2 gaps vs LS-Oracle
    gap_ls_oracle      = 0.0
    gap_drl_real       = _clean(g("DRL Real Gap vs LS-Oracle (%)"))
    gap_rh_la_real     = _clean(_avg_gap(results_df, "LS-Oracle Reward", "LS-Oracle Valid",
                                          "RH-Lookahead Real Reward", "RH-Lookahead Real Valid"))
    gap_rh_greedy_real = _clean(_avg_gap(results_df, "LS-Oracle Reward", "LS-Oracle Valid",
                                          "RH-Greedy Real Reward",    "RH-Greedy Real Valid"))
    gap_mc_rollout     = _clean(_avg_gap(results_df, "LS-Oracle Reward", "LS-Oracle Valid",
                                          "MC-Rollout Reward",         "MC-Rollout Valid"))

    # Inference times
    inf_ls_exact      = _clean(g("LS-Exact Avg Inference (ms)"))
    inf_drl_det       = _clean(g("DRL Real Avg Inference (ms)"))   # same model
    inf_drl_real      = _clean(g("DRL Real Avg Inference (ms)"))
    inf_hga_lns       = _clean(g("HGA-LNS Avg Inference (ms)"))
    inf_rh_lookahead  = _clean(g("RH-Lookahead Avg Inference (ms)"))
    inf_rh_greedy     = _clean(g("RH-Greedy Avg Inference (ms)"))
    inf_ls_oracle     = _clean(g("LS-Oracle Avg Inference (ms)"))
    inf_rh_la_real    = _clean(g("RH-Lookahead Real Avg Inference (ms)"))
    inf_mc_rollout    = _clean(g("MC-Rollout Avg Inference (ms)"))

    row = 1

    # ── Block 1 — Deterministic ───────────────────────────────
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    _header_cell(ws.cell(row, 1), "BLOQUE 1 — MUNDO DETERMINISTA", _BLUE_DARK)
    row += 1

    for c, h in enumerate(
        ["Método", "Avg Reward", "Gap vs LS-Exact (%)", "Avg Inference (ms)"], 1
    ):
        _header_cell(ws.cell(row, c), h, _BLUE_LIGHT, font_color="000000")
    row += 1

    b1 = [
        ("LS-Exact",     _clean(g("LS-Exact Avg Reward")),    gap_ls_exact,     inf_ls_exact),
        ("DRL Det",      _clean(g("DRL Det Avg Reward")),      gap_drl_det,      inf_drl_det),
        ("HGA-LNS",      _clean(g("HGA-LNS Avg Reward")),      gap_hga_lns,      inf_hga_lns),
        ("RH-Lookahead", _clean(g("RH-Lookahead Avg Reward")), gap_rh_lookahead, inf_rh_lookahead),
        ("RH-Greedy",    _clean(g("RH-Greedy Avg Reward")),    gap_rh_greedy,    inf_rh_greedy),
    ]
    for name, reward, gap, inf_ms in b1:
        _data_cell(ws.cell(row, 1), name)
        _data_cell(ws.cell(row, 2), reward,  '#,##0')
        _data_cell(ws.cell(row, 3), gap,     '0.00"%"', fill_hex=_YELLOW)
        _data_cell(ws.cell(row, 4), inf_ms,  '0.00')
        row += 1

    row += 1  # blank separator

    # ── Block 2 — Stochastic ──────────────────────────────────
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    _header_cell(ws.cell(row, 1), "BLOQUE 2 — MUNDO ESTOCÁSTICO", _PURPLE_DARK)
    row += 1

    for c, h in enumerate(
        ["Método", "Avg Reward", "Gap vs LS-Oracle (%)", "Avg Inference (ms)"], 1
    ):
        _header_cell(ws.cell(row, c), h, _PURPLE_LIGHT, font_color="000000")
    row += 1

    b2 = [
        ("LS-Oracle",         _clean(g("LS-Oracle Avg Reward")),            gap_ls_oracle,      inf_ls_oracle),
        ("DRL Real",          _clean(g("DRL Real Avg Reward")),             gap_drl_real,       inf_drl_real),
        ("RH-Lookahead Real", _clean(g("RH-Lookahead Real Avg Reward")),    gap_rh_la_real,     inf_rh_la_real),
        ("RH-Greedy Real",    _clean(g("RH-Greedy Real Avg Reward")),       gap_rh_greedy_real, None),
        ("MC-Rollout",        _clean(g("MC-Rollout Avg Reward")),           gap_mc_rollout,     inf_mc_rollout),
    ]
    for name, reward, gap, inf_ms in b2:
        _data_cell(ws.cell(row, 1), name)
        _data_cell(ws.cell(row, 2), reward,  '#,##0')
        _data_cell(ws.cell(row, 3), gap,     '0.00"%"', fill_hex=_YELLOW)
        _data_cell(ws.cell(row, 4), inf_ms,  '0.00')
        row += 1

    row += 1  # blank separator

    # ── Block 3 — Timing ──────────────────────────────────────
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    _header_cell(ws.cell(row, 1), "BLOQUE 3 — TIEMPO DE INFERENCIA Y ENTRENAMIENTO", _GREEN_DARK)
    row += 1

    for c, h in enumerate(["Métrica", "Valor (ms / s)"], 1):
        _header_cell(ws.cell(row, c), h, _GREEN_LIGHT, font_color="000000")
    row += 1

    b3 = [
        ("DRL Training Time (s)",                _clean(g("DRL Training Time (s)"))),
        ("DRL Real Avg Inference (ms)",           inf_drl_real),
        ("HGA-LNS Avg Inference (ms)",            inf_hga_lns),
        ("LS-Exact Avg Inference (ms)",           inf_ls_exact),
        ("RH-Greedy Avg Inference (ms)",          inf_rh_greedy),
        ("RH-Lookahead Avg Inference (ms)",       inf_rh_lookahead),
        ("RH-Lookahead Real Avg Inference (ms)",  inf_rh_la_real),
        ("MC-Rollout Avg Inference (ms)",         inf_mc_rollout),
        ("LS-Oracle Avg Inference (ms)",          inf_ls_oracle),
    ]
    for metric, value in b3:
        _data_cell(ws.cell(row, 1), metric)
        _data_cell(ws.cell(row, 2), value, '0.00')
        row += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 22
    ws.freeze_panes = "A4"


# ═════════════════════════════════════════════════════════════
#  PER NODE RESULTS SHEET
# ═════════════════════════════════════════════════════════════

# Method definitions: (display_name, has_status_col, df_prefix)
_DET_METHODS = [
    ("LS-Exact",     True,  "LS-Exact"),
    ("DRL Det",      False, "DRL Det"),
    ("HGA-LNS",      True,  "HGA-LNS"),
    ("RH-Lookahead", False, "RH-Lookahead"),
    ("RH-Greedy",    False, "RH-Greedy"),
]
_STOCH_METHODS = [
    ("LS-Oracle",         True,  "LS-Oracle"),
    ("DRL Real",          False, "DRL Real"),
    ("RH-Lookahead Real", False, "RH-Lookahead Real"),
    ("RH-Greedy Real",    False, "RH-Greedy Real"),
    ("MC-Rollout",        False, "MC-Rollout"),
]

def _method_subcols(has_status):
    return ["Status", "Route", "Reward", "Duration", "Valid"] if has_status \
           else ["Route", "Reward", "Duration", "Valid"]


def _build_col_info(methods, block_fill):
    """Return list of (row3_label, df_col, num_format, fill_hex) for data cols."""
    cols = []
    for _, has_status, prefix in methods:
        for sub in _method_subcols(has_status):
            df_col = f"{prefix} {sub}"
            nfmt = '#,##0' if sub == "Reward" else ('0.00' if sub == "Duration" else None)
            cols.append((sub, df_col, nfmt, block_fill))
    return cols


def _row_gap(row_dict, bench_col, bench_valid_col, method_prefix):
    b_valid = bool(row_dict.get(bench_valid_col, False))
    v_valid = bool(row_dict.get(f"{method_prefix} Valid", False))
    if not b_valid or not v_valid:
        return None
    g = _compute_gap(row_dict.get(bench_col), row_dict.get(f"{method_prefix} Reward"))
    return None if (g is None or np.isnan(g)) else g


def _build_per_node(wb, results_df):
    ws = wb.create_sheet("Per Node Results")

    fixed_cols = [
        ("Start Node",      "Start Node",      None,    _WHITE),
        ("Eval Day Index",  "Eval Day Index",  None,    _WHITE),
        ("Abs Day Index",   "Abs Day Index",   None,    _WHITE),
    ]

    det_data_cols   = _build_col_info(_DET_METHODS,   _BLUE_LIGHT)
    stoch_data_cols = _build_col_info(_STOCH_METHODS, _PURPLE_LIGHT)

    det_gap_cols = [
        (f"Gap {disp} (%)", f"__gap_det_{pfx}", '0.00"%"', _YELLOW)
        for disp, _, pfx in _DET_METHODS
    ]
    stoch_gap_cols = [
        (f"Gap {disp} (%)", f"__gap_stoch_{pfx}", '0.00"%"', _YELLOW)
        for disp, _, pfx in _STOCH_METHODS
    ]

    # full column list in sheet order
    col_info = fixed_cols + det_data_cols + det_gap_cols + stoch_data_cols + stoch_gap_cols

    n_fixed      = len(fixed_cols)
    n_det_data   = len(det_data_cols)
    n_det_gap    = len(det_gap_cols)
    n_stoch_data = len(stoch_data_cols)
    n_stoch_gap  = len(stoch_gap_cols)

    det_span   = n_det_data + n_det_gap
    stoch_span = n_stoch_data + n_stoch_gap

    det_start   = n_fixed + 1
    det_end     = n_fixed + det_span
    stoch_start = det_end + 1
    stoch_end   = stoch_start + stoch_span - 1

    # ── Row 1: super-headers ──────────────────────────────────
    for i, (label, *_) in enumerate(fixed_cols, 1):
        c = ws.cell(1, i)
        c.value = label
        c.fill  = _fill(_WHITE)
        c.font  = _font(bold=True)
        c.alignment = _center()
        c.border = _border()

    ws.merge_cells(start_row=1, start_column=det_start, end_row=1, end_column=det_end)
    _header_cell(ws.cell(1, det_start), "BLOQUE 1 — MUNDO DETERMINISTA", _BLUE_DARK)

    ws.merge_cells(start_row=1, start_column=stoch_start, end_row=1, end_column=stoch_end)
    _header_cell(ws.cell(1, stoch_start), "BLOQUE 2 — MUNDO ESTOCÁSTICO", _PURPLE_DARK)

    # ── Row 2: method sub-headers ─────────────────────────────
    for i in range(1, n_fixed + 1):
        ws.cell(2, i).fill   = _fill(_WHITE)
        ws.cell(2, i).border = _border()

    col_ptr = n_fixed + 1
    for disp, has_status, _ in _DET_METHODS:
        span = len(_method_subcols(has_status))
        end  = col_ptr + span - 1
        if span > 1:
            ws.merge_cells(start_row=2, start_column=col_ptr, end_row=2, end_column=end)
        _header_cell(ws.cell(2, col_ptr), disp, _BLUE_LIGHT, font_color="000000")
        col_ptr = end + 1

    gap_det_end = col_ptr + n_det_gap - 1
    if n_det_gap > 1:
        ws.merge_cells(start_row=2, start_column=col_ptr, end_row=2, end_column=gap_det_end)
    _header_cell(ws.cell(2, col_ptr), "Gaps vs LS-Exact", _YELLOW, font_color="000000")
    col_ptr = gap_det_end + 1

    for disp, has_status, _ in _STOCH_METHODS:
        span = len(_method_subcols(has_status))
        end  = col_ptr + span - 1
        if span > 1:
            ws.merge_cells(start_row=2, start_column=col_ptr, end_row=2, end_column=end)
        _header_cell(ws.cell(2, col_ptr), disp, _PURPLE_LIGHT, font_color="000000")
        col_ptr = end + 1

    gap_stoch_end = col_ptr + n_stoch_gap - 1
    if n_stoch_gap > 1:
        ws.merge_cells(start_row=2, start_column=col_ptr, end_row=2, end_column=gap_stoch_end)
    _header_cell(ws.cell(2, col_ptr), "Gaps vs LS-Oracle", _YELLOW, font_color="000000")

    # ── Row 3: column names ───────────────────────────────────
    for c, (label, _, _, fill) in enumerate(col_info, 1):
        cell = ws.cell(3, c)
        cell.value = label
        cell.fill  = _fill(fill)
        cell.font  = _font(bold=True)
        cell.alignment = _center()
        cell.border = _border()

    # ── Data rows ─────────────────────────────────────────────
    for r_idx, (_, row_series) in enumerate(results_df.iterrows(), 4):
        rd = row_series.to_dict()

        det_gaps = {
            pfx: _row_gap(rd, "LS-Exact Reward",  "LS-Exact Valid",  pfx)
            for _, _, pfx in _DET_METHODS
        }
        stoch_gaps = {
            pfx: _row_gap(rd, "LS-Oracle Reward", "LS-Oracle Valid", pfx)
            for _, _, pfx in _STOCH_METHODS
        }

        for c, (label, df_col, nfmt, fill) in enumerate(col_info, 1):
            cell = ws.cell(r_idx, c)
            cell.fill      = _fill(fill)
            cell.alignment = _left()
            cell.border    = _border()

            if df_col.startswith("__gap_det_"):
                pfx = df_col[len("__gap_det_"):]
                val = det_gaps.get(pfx)
                cell.value = val
                if val is not None and nfmt:
                    cell.number_format = nfmt

            elif df_col.startswith("__gap_stoch_"):
                pfx = df_col[len("__gap_stoch_"):]
                val = stoch_gaps.get(pfx)
                cell.value = val
                if val is not None and nfmt:
                    cell.number_format = nfmt

            else:
                raw = rd.get(df_col)
                if isinstance(raw, list):
                    val = str(raw)
                elif df_col.endswith((" Reward", " Duration")):
                    val = _clean(raw)
                else:
                    val = raw
                cell.value = val
                if val is not None and nfmt:
                    cell.number_format = nfmt

    # ── Column widths ─────────────────────────────────────────
    for c, (label, *_) in enumerate(col_info, 1):
        if "Route" in label:
            w = 32
        elif "Gap" in label:
            w = 16
        elif label in ("Status",):
            w = 14
        elif label in ("Start Node", "Eval Day Index", "Abs Day Index"):
            w = 14
        else:
            w = 13
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.freeze_panes = "D4"
